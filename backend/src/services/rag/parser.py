from typing import Optional
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from src.utils.logger import logger


class ParseResult:
    def __init__(self, text: str, pages: Optional[list[str]] = None, sections: Optional[list[dict]] = None):
        self.text = text
        self.pages = pages or []
        self.sections = sections or []


class DocumentParser:
    SUPPORTED_FORMATS = {".pdf", ".docx", ".pptx", ".txt", ".md"}

    def parse(self, file_path: str, file_type: str) -> ParseResult:
        ext = f".{file_type.lower()}"

        if file_type == "pdf" or ext == ".pdf":
            return self._parse_pdf(file_path)
        elif file_type == "docx" or ext == ".docx":
            return self._parse_docx(file_path)
        elif file_type == "pptx" or ext == ".pptx":
            return self._parse_pptx(file_path)
        elif file_type in ("txt", "md") or ext in (".txt", ".md"):
            return self._parse_text(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    def _parse_pdf(self, file_path: str) -> ParseResult:
        pages = []
        try:
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                for page in reader.pages:
                    text = page.extract_text() or ""
                    pages.append(text)
        except Exception as e:
            logger.warning(f"pypdf failed on {file_path}: {e}")

        combined = "\n\n".join(pages).strip()
        if not combined:
            try:
                import fitz
                doc = fitz.open(file_path)
                pages = [page.get_text() for page in doc]
                doc.close()
                combined = "\n\n".join(pages).strip()
            except Exception as e2:
                logger.warning(f"fitz fallback failed on {file_path}: {e2}")

        return ParseResult(
            text=combined,
            pages=pages,
        )

    def _parse_docx(self, file_path: str) -> ParseResult:
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)

        sections = []
        for i, p in enumerate(doc.paragraphs):
            if p.style and "heading" in (p.style.name or "").lower():
                sections.append({"heading": p.text, "index": i})

        return ParseResult(
            text=text,
            pages=[text],
            sections=sections,
        )

    def _parse_pptx(self, file_path: str) -> ParseResult:
        prs = Presentation(file_path)
        slides = []
        for slide in prs.slides:
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            slides.append("\n".join(slide_text))

        return ParseResult(
            text="\n\n".join(slides),
            pages=slides,
        )

    def _parse_text(self, file_path: str) -> ParseResult:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

        return ParseResult(
            text=text,
            pages=[text],
        )


document_parser = DocumentParser()
